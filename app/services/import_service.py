"""Bulk product import from an Excel workbook, optionally with a ZIP of product images."""

from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import PurePosixPath

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError as PydanticValidationError
from starlette.concurrency import run_in_threadpool

from app.core.exceptions import (
    DataAccessError,
    DuplicateProductError,
    ProductNotFoundError,
    ValidationError,
)
from app.models.schemas import ProductCreate, ProductUpdate
from app.services.image_service import ImageService
from app.services.product_service import ProductService

logger = logging.getLogger(__name__)

MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_ZIP_BYTES = 40 * 1024 * 1024
MAX_ROWS = 200
MAX_ZIP_IMAGES = 400
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_CELL_LENGTH = 5000

TEMPLATE_FILENAME = "nawa-product-import-template.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNS = (
    ("Name", "Required. Product name, up to 255 characters.", 34),
    ("Category", "Required. Pick one from the list or type a new category.", 22),
    ("Description", "Optional. Shown on the product page.", 48),
    (
        "Specifications",
        "Optional. One per line as Label: Value (press Alt+Enter for a new line), or separate with ;",
        44,
    ),
    ("Keywords", "Optional. Words customers may search, separated by commas. Not shown on the site.", 30),
    ("Featured", "Optional. Yes to show on the home page. Leave empty for No.", 12),
    (
        "Images",
        "Optional. File names from your ZIP, separated by commas. The first one is the main image.",
        40,
    ),
)
HEADERS = tuple(name for name, _, _ in COLUMNS)
_TRUE_VALUES = {"yes", "y", "true", "1", "x", "featured"}


class ImportFileError(Exception):
    """The upload cannot be read at all (as opposed to a single bad row)."""


@dataclass
class RowResult:
    row: int
    name: str
    status: str  # added, updated, skipped, error
    notes: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    rows: list[RowResult] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for row in self.rows if row.status == status)


class MemoryUpload:
    """Minimal UploadFile stand-in so ZIP members reuse the normal image pipeline."""

    def __init__(self, filename: str, content: bytes) -> None:
        self.filename = filename
        self._content = content

    async def read(self, size: int = -1) -> bytes:
        return self._content if size < 0 else self._content[:size]


class ImageArchive:
    """Safe, read-only view of a ZIP of images, matched by file name (folders ignored)."""

    def __init__(self, data: bytes | None, max_image_bytes: int) -> None:
        self.max_image_bytes = max_image_bytes
        self._members: dict[str, zipfile.ZipInfo] = {}
        self._zip: zipfile.ZipFile | None = None
        if not data:
            return
        if len(data) > MAX_ZIP_BYTES:
            raise ImportFileError(
                f"The images ZIP is larger than {MAX_ZIP_BYTES // (1024 * 1024)} MB."
            )
        try:
            archive = zipfile.ZipFile(BytesIO(data))
            infos = archive.infolist()
        except zipfile.BadZipFile as exc:
            raise ImportFileError("The images file is not a valid ZIP archive.") from exc
        total = 0
        for info in infos:
            path = PurePosixPath(info.filename)
            name = path.name.lower()
            if info.is_dir() or not name or name.startswith("._") or "__macosx" in info.filename.lower():
                continue
            if path.suffix.lower() not in ImageService.allowed_extensions:
                continue
            if info.flag_bits & 0x1:
                raise ImportFileError("Password-protected ZIP files are not supported.")
            total += info.file_size
            if total > MAX_UNCOMPRESSED_BYTES or len(self._members) >= MAX_ZIP_IMAGES:
                raise ImportFileError(
                    f"The ZIP has too many or too large images (limit {MAX_ZIP_IMAGES} images)."
                )
            self._members.setdefault(name, info)
        self._zip = archive

    def find(self, filename: str) -> str | None:
        key = PurePosixPath(filename.replace("\\", "/")).name.lower()
        return key if key in self._members else None

    def read(self, key: str) -> MemoryUpload:
        info = self._members[key]
        if info.file_size > self.max_image_bytes:
            raise ValidationError(f"Image '{info.filename}' exceeds the size limit.")
        assert self._zip is not None
        with self._zip.open(info) as handle:
            return MemoryUpload(key, handle.read(self.max_image_bytes + 1))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()[:MAX_CELL_LENGTH]


def _split(value: str, separators: str) -> list[str]:
    return [part.strip() for part in re.split(separators, value) if part.strip()]


def read_workbook_rows(data: bytes) -> list[tuple[int, dict[str, str]]]:
    """Return (sheet row number, values by column) for each non-empty product row."""
    if not data:
        raise ImportFileError("Choose the filled Excel file (.xlsx) to import.")
    if len(data) > MAX_WORKBOOK_BYTES:
        raise ImportFileError(
            f"The Excel file is larger than {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB."
        )
    try:
        archive = zipfile.ZipFile(BytesIO(data))
        names = archive.namelist()
        if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES // 4:
            raise ImportFileError("The Excel file is too large once unpacked.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("This is not an .xlsx Excel file. Use the template provided.") from exc
    if "xl/workbook.xml" not in names:
        raise ImportFileError("This is not an .xlsx Excel file. Use the template provided.")
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise ImportFileError("Excel files with macros are not accepted. Save as .xlsx.")

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = workbook["Products"] if "Products" in workbook.sheetnames else workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ImportFileError("The Products sheet is empty.")
        header = {_cell_text(cell).lower(): index for index, cell in enumerate(header_row) if cell}
        missing = [name for name in ("Name", "Category") if name.lower() not in header]
        if missing:
            raise ImportFileError(
                "The first row must contain the column titles from the template (missing: "
                + ", ".join(missing)
                + ")."
            )
        rows: list[tuple[int, dict[str, str]]] = []
        for number, values in enumerate(iterator, start=2):
            record = {
                title: _cell_text(values[header[title.lower()]]) if header.get(title.lower(), len(values)) < len(values) else ""
                for title in HEADERS
            }
            if not any(record.values()):
                continue
            if len(rows) >= MAX_ROWS:
                raise ImportFileError(f"Too many rows. Import up to {MAX_ROWS} products at a time.")
            rows.append((number, record))
        workbook.close()
    except ImportFileError:
        raise
    except Exception as exc:  # openpyxl raises many types for damaged files
        logger.warning("Could not read product workbook: %s", exc)
        raise ImportFileError("The Excel file could not be read. Re-save it as .xlsx and try again.") from exc
    return rows


def build_template(categories: list[str]) -> bytes:
    """Create the workbook clients fill in: Products (empty), Instructions, Example."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    header_fill = PatternFill("solid", fgColor="0A1A14")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, (title, _, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = header_fill
        cell.font = header_font
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=len(COLUMNS)):
        for cell in row:
            cell.alignment = wrap

    featured = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    featured.add(f"F2:F{MAX_ROWS + 1}")
    sheet.add_data_validation(featured)

    lists = workbook.create_sheet("Lists")
    for index, name in enumerate(categories, start=1):
        lists.cell(row=index, column=1, value=name)
    if categories:
        # Suggestions only: new categories can still be typed in.
        choice = DataValidation(
            type="list",
            formula1=f"Lists!$A$1:$A${len(categories)}",
            allow_blank=True,
            showErrorMessage=False,
        )
        choice.add(f"B2:B{MAX_ROWS + 1}")
        sheet.add_data_validation(choice)
    lists.sheet_state = "hidden"

    guide = workbook.create_sheet("Instructions", 1)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 92
    guide["A1"] = "How to fill in the Products sheet"
    guide["A1"].font = Font(bold=True, size=14)
    for index, (title, text, _) in enumerate(COLUMNS, start=3):
        guide.cell(row=index, column=1, value=title).font = Font(bold=True)
        guide.cell(row=index, column=2, value=text).alignment = wrap
    tips = [
        "One product per row. Leave the first row (titles) as it is. Keep the Products sheet name.",
        f"Up to {MAX_ROWS} products per upload.",
        "Images: put all picture files (JPG, PNG or WEBP) in one ZIP file and upload it together with the Excel file.",
        "In the Images column write the file names exactly as they are in the ZIP, for example gate-valve-1.jpg, gate-valve-2.jpg",
        "The same product can have several images. The first file name becomes the main image.",
        "A product with the same name and category that already exists is skipped, unless you tick the update option.",
        "The Example sheet shows two filled rows. Do not copy it back with the examples in it.",
    ]
    guide.cell(row=len(COLUMNS) + 4, column=1, value="Good to know").font = Font(bold=True)
    for offset, tip in enumerate(tips, start=len(COLUMNS) + 5):
        guide.cell(row=offset, column=1, value="•").alignment = Alignment(horizontal="right")
        guide.cell(row=offset, column=2, value=tip).alignment = wrap

    example = workbook.create_sheet("Example", 2)
    for index, (title, _, width) in enumerate(COLUMNS, start=1):
        cell = example.cell(row=1, column=index, value=title)
        cell.fill = header_fill
        cell.font = header_font
        example.column_dimensions[get_column_letter(index)].width = width
    samples = [
        (
            "Gate valve 2 in",
            "Valves",
            "Carbon steel gate valve for water, oil and gas lines.",
            "Size: 2 in\nPressure: Class 150\nStandard: API 6D",
            "gate valve, flow control, valve",
            "Yes",
            "gate-valve-1.jpg, gate-valve-2.jpg",
        ),
        (
            "Nitrile safety gloves",
            "Safety",
            "Oil-resistant gloves for workshop and field use.",
            "Material: Nitrile; Sizes: S, M, L, XL",
            "gloves, ppe, hand protection",
            "",
            "gloves.jpg",
        ),
    ]
    for row_index, sample in enumerate(samples, start=2):
        for column, value in enumerate(sample, start=1):
            example.cell(row=row_index, column=column, value=value).alignment = wrap

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _validation_text(exc: PydanticValidationError) -> str:
    first = exc.errors()[0]
    field_name = str(first["loc"][-1]).replace("_", " ").title() if first.get("loc") else "Row"
    message = str(first.get("msg", "invalid value")).removeprefix("Value error, ")
    return f"{field_name}: {message}"


class ProductImportService:
    def __init__(self, product_service: ProductService, image_service: ImageService) -> None:
        self.products = product_service
        self.images = image_service

    async def import_products(
        self, workbook: bytes, archive: bytes | None, update_existing: bool
    ) -> ImportReport:
        rows = await run_in_threadpool(read_workbook_rows, workbook)
        if not rows:
            raise ImportFileError("No product rows found. Fill in the Products sheet and try again.")
        images = ImageArchive(archive, self.images.max_upload_size_bytes)
        report = ImportReport()
        for number, record in rows:
            report.rows.append(await self._import_row(number, record, images, update_existing))
        return report

    async def _import_row(
        self, number: int, record: dict[str, str], archive: ImageArchive, update_existing: bool
    ) -> RowResult:
        result = RowResult(row=number, name=record["Name"] or "(no name)", status="error")
        try:
            fields = {
                "name": record["Name"],
                "category": record["Category"],
                "description": record["Description"],
                "specs": _split(record["Specifications"], r"[\n;]"),
                "tags": record["Keywords"],
                "featured": record["Featured"].strip().lower() in _TRUE_VALUES,
            }
            payload = ProductCreate(**fields)
        except PydanticValidationError as exc:
            result.notes.append(_validation_text(exc))
            return result

        uploads: list[MemoryUpload] = []
        for filename in _split(record["Images"], r"[,;\n]"):
            key = archive.find(filename)
            if key is None:
                result.notes.append(f"Image not found in the ZIP: {filename}")
                continue
            try:
                uploads.append(archive.read(key))
            except ValidationError as exc:
                result.notes.append(str(exc))

        try:
            await self.products.create_product(payload, uploads)
            result.status = "added"
        except DuplicateProductError:
            if not update_existing:
                result.status = "skipped"
                result.notes.append("Already exists (same name and category).")
                return result
            result.status = await self._update_existing(payload, uploads, result)
        except (ValidationError, ProductNotFoundError) as exc:
            result.notes.append(str(exc))
        except DataAccessError:
            result.notes.append("Could not save this row. Try again.")
        return result

    async def _update_existing(
        self, payload: ProductCreate, uploads: list[MemoryUpload], result: RowResult
    ) -> str:
        product_id = self.products.find_product_id(payload.name, payload.category)
        if product_id is None:
            result.notes.append("Could not find the existing product to update.")
            return "error"
        try:
            await self.products.update_product(
                product_id, ProductUpdate(**payload.model_dump()), uploads
            )
        except (ValidationError, ProductNotFoundError, DuplicateProductError) as exc:
            result.notes.append(str(exc))
            return "error"
        except DataAccessError:
            result.notes.append("Could not update this product. Try again.")
            return "error"
        if uploads:
            result.notes.append("New images were added to the existing ones.")
        return "updated"
