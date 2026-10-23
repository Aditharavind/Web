"""Bulk product import from Excel with a ZIP of images."""

import io
import zipfile

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.core.config import get_settings
from app.main import app
from app.models.schemas import ProductCreate
from app.services import import_service
from app.services.dependencies import get_import_service, get_product_service
from app.services.import_service import (
    HEADERS,
    ImportFileError,
    build_template,
    read_workbook_rows,
)
from tests.test_admin_routes import login


def workbook_bytes(rows, headers=HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.title = "Products"
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def zip_bytes(names):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name in names:
            image = io.BytesIO()
            Image.new("RGB", (30, 30), "blue").save(image, "JPEG")
            archive.writestr(name, image.getvalue())
    return output.getvalue()


def row(name, category="Valves", description="", specs="", tags="", featured="", images=""):
    return [name, category, description, specs, tags, featured, images]


def test_template_has_expected_sheets_and_headers():
    book = load_workbook(io.BytesIO(build_template(["Valves", "Safety"])))
    assert book.sheetnames == ["Products", "Instructions", "Example", "Lists"]
    assert [cell.value for cell in book["Products"][1]] == list(HEADERS)
    assert book["Lists"].sheet_state == "hidden"
    # The template itself must not contain importable sample rows.
    assert read_workbook_rows(build_template([])) == []


@pytest.mark.asyncio
async def test_import_creates_products_with_multiple_images_and_reports_rows():
    report = await get_import_service().import_products(
        workbook_bytes(
            [
                row("Gate valve", "Valves", "Steel valve", "Size: 2 in\nClass: 150", "gate, flow", "Yes",
                    "a.jpg, b.JPG ,missing.jpg"),
                row("Bad row", ""),
                row("Gloves", "Safety", specs="Material: Nitrile; Sizes: S, M", images="folder/c.jpg"),
                [None] * 7,
            ]
        ),
        zip_bytes(["a.jpg", "photos/b.jpg", "c.jpg", "__MACOSX/._a.jpg", "notes.txt"]),
        update_existing=False,
    )
    assert [(r.row, r.status) for r in report.rows] == [(2, "added"), (3, "error"), (4, "added")]
    assert report.rows[0].notes == ["Image not found in the ZIP: missing.jpg"]
    assert "Category" in report.rows[1].notes[0]

    service = get_product_service()
    valve = service.search_catalog(1, "Gate valve").items[0]
    assert valve.featured is True
    assert valve.specs == ["Size: 2 in", "Class: 150"]
    assert len(valve.images) == 2
    gloves = service.search_catalog(1, "Gloves").items[0]
    assert gloves.specs == ["Material: Nitrile", "Sizes: S, M"]
    assert len(gloves.images) == 1


@pytest.mark.asyncio
async def test_duplicates_are_skipped_or_updated_with_extra_images():
    service = get_product_service()
    first = await service.create_product(ProductCreate(name="Seal", category="Parts"), [])
    sheet = workbook_bytes([row("seal", "parts", "New text", images="x.jpg")])
    archive = zip_bytes(["x.jpg"])

    skipped = await get_import_service().import_products(sheet, archive, update_existing=False)
    assert skipped.rows[0].status == "skipped"
    assert service.get_product(first.id).description == ""

    updated = await get_import_service().import_products(sheet, archive, update_existing=True)
    assert updated.rows[0].status == "updated"
    product = service.get_product(first.id)
    assert product.description == "New text"
    assert len(product.images) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "data,message",
    [
        (b"", "Choose the filled Excel file"),
        (b"not a workbook", "not an .xlsx"),
        (workbook_bytes([], headers=("Title", "Type")), "missing: Name, Category"),
    ],
)
async def test_unreadable_workbooks_are_rejected(data, message):
    with pytest.raises(ImportFileError, match=message):
        await get_import_service().import_products(data, None, False)


@pytest.mark.asyncio
async def test_limits_and_unsafe_archives_are_rejected(monkeypatch):
    many = workbook_bytes([row(f"P{i}") for i in range(3)])
    monkeypatch.setattr(import_service, "MAX_ROWS", 2)
    with pytest.raises(ImportFileError, match="Too many rows"):
        await get_import_service().import_products(many, None, False)
    with pytest.raises(ImportFileError, match="not a valid ZIP"):
        await get_import_service().import_products(workbook_bytes([row("A")]), b"nope", False)
    macro = io.BytesIO()
    with zipfile.ZipFile(macro, "w") as archive:
        archive.writestr("xl/workbook.xml", "<x/>")
        archive.writestr("xl/vbaProject.bin", "x")
    with pytest.raises(ImportFileError, match="macros"):
        read_workbook_rows(macro.getvalue())


@pytest.mark.asyncio
async def test_admin_import_endpoint_requires_login_csrf_and_shows_report():
    prefix = get_settings().admin_route
    sheet = workbook_bytes([row("Ball valve", "Valves", images="v.jpg")])
    files = {"workbook": ("products.xlsx", sheet), "images_zip": ("pics.zip", zip_bytes(["v.jpg"]))}
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver") as client:
        anonymous = await client.post(f"{prefix}/products/import", data={"csrf_token": "x"}, files=files)
        assert anonymous.status_code == 303
        assert anonymous.headers["location"] == f"{prefix}/login"
        assert (await client.get(f"{prefix}/products/import-template.xlsx")).status_code == 303

        token = await login(client)
        forged = await client.post(f"{prefix}/products/import", data={"csrf_token": "bad"}, files=files)
        assert forged.status_code == 403

        template = await client.get(f"{prefix}/products/import-template.xlsx")
        assert template.status_code == 200
        assert "attachment" in template.headers["content-disposition"]
        assert load_workbook(io.BytesIO(template.content)).sheetnames[0] == "Products"

        done = await client.post(f"{prefix}/products/import", data={"csrf_token": token}, files=files)
        assert done.status_code == 200
        assert "Ball valve" in done.text and "Added" in done.text
        page = await client.get(f"{prefix}/products")
        assert "Import from Excel" in page.text and "Ball valve" in page.text

        broken = await client.post(
            f"{prefix}/products/import",
            data={"csrf_token": token},
            files={"workbook": ("x.xlsx", b"junk")},
        )
        assert broken.status_code == 303 and "error=" in broken.headers["location"]


@pytest.mark.asyncio
async def test_too_many_image_references_are_rejected(monkeypatch):
    monkeypatch.setattr(import_service, "MAX_IMAGE_REFERENCES", 3)
    sheet = workbook_bytes([row("A", images="a.jpg, b.jpg"), row("B", images="a.jpg, b.jpg")])
    with pytest.raises(ImportFileError, match="Too many image names"):
        await get_import_service().import_products(sheet, zip_bytes(["a.jpg", "b.jpg"]), False)
